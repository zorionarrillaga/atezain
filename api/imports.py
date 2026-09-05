"""Bounded spreadsheet ingestion. Reject uncertain money and report every skipped row."""
from __future__ import annotations

import csv
import datetime
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from defusedxml.ElementTree import iterparse
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from api.schemas import RejectedRow, UploadOut

MAX_UPLOAD_BYTES = 1024 * 1024
MAX_ROWS = 500
COLUMNS = ("id", "customer", "amount", "currency", "issued", "due", "status")
EXTRAS = ("contact", "note", "email_subject", "email_body")
STATUSES = {"open", "paid", "reminded", "promised", "disputed", "cancelled"}
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}", re.ASCII)
SLASHED_DATE = re.compile(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", re.ASCII)


def parse_rows(raw: bytes, filename: str) -> list[dict]:
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "file exceeds the upload byte limit")
    workbook = None
    try:
        if filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 128 or sum(m.file_size for m in members) > 8 * MAX_UPLOAD_BYTES:
                    raise HTTPException(413, "expanded workbook exceeds the safety limit")
                if any(m.flag_bits & 1 or (m.file_size > 100_000 and m.file_size > 200 * max(m.compress_size, 1)) for m in members):
                    raise HTTPException(413, "encrypted or excessively compressed workbook")
                nodes = 0
                for member in members:
                    if not member.filename.lower().endswith((".xml", ".rels")):
                        continue
                    depth = 0
                    for event, element in iterparse(io.BytesIO(archive.read(member)), events=("start", "end"),
                                                    forbid_dtd=True, forbid_entities=True, forbid_external=True):
                        if event == "start":
                            nodes += 1
                            depth += 1
                            if nodes > 100000 or depth > 64:
                                raise HTTPException(413, "workbook XML structure exceeds the safety limit")
                        else:
                            depth -= 1
                            element.clear()
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheet = workbook.active
            if sheet is None:
                raise ValueError("missing worksheet")
            if sheet.max_column and sheet.max_column > len(COLUMNS + EXTRAS):
                raise HTTPException(422, "workbook has unexpected columns")

            def cells():
                for row in sheet.iter_rows(max_row=MAX_ROWS + 2):
                    values = []
                    for cell in row:
                        if cell.data_type == "f":
                            raise HTTPException(422, "formula cells are not accepted; export values only")
                        v = cell.value
                        if isinstance(v, (datetime.datetime, datetime.date)):
                            v = v.date().isoformat() if isinstance(v, datetime.datetime) else v.isoformat()
                        values.append("" if v is None else str(v))
                    yield values
            iterator = iter(cells())
        elif filename.lower().endswith(".csv"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError as e:
                raise HTTPException(415, "the file is not UTF-8 text or .xlsx") from e
            header_line = text.splitlines()[0] if text else ""
            delimiter = ";" if ";" in header_line and "," not in header_line else ","
            iterator = iter(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
        else:
            raise HTTPException(415, "upload a CSV or XLSX file")
        header = [v.strip().lower() for v in next(iterator, [])]
        if not header or len(header) != len(set(header)) or any(k not in COLUMNS + EXTRAS for k in header):
            raise HTTPException(422, "missing, duplicate, blank or unknown column headers")
        if not {"id", "customer", "amount", "issued", "due"}.issubset(header):
            raise HTTPException(422, "required columns: id, customer, amount, issued, due")
        rows = []
        for values in iterator:
            if not any(v.strip() for v in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise HTTPException(413, f"more than {MAX_ROWS} rows")
            if len(values) != len(header):
                raise HTTPException(422, f"row {len(rows) + 2} has a different number of cells than the header")
            if any(len(v) > 8000 or "\x00" in v for v in values):
                raise HTTPException(422, f"row {len(rows) + 2} contains an oversized cell or a NUL character")
            rows.append(dict(zip(header, values)))
        if not rows:
            raise HTTPException(422, "the file contains no invoice rows")
        return rows
    except HTTPException:
        raise
    except (ValueError, KeyError, OSError, zipfile.BadZipFile, csv.Error, StopIteration, DefusedXmlException, ParseError) as e:
        raise HTTPException(422, "the spreadsheet is malformed; export a fresh CSV or XLSX") from e
    finally:
        if workbook is not None:
            workbook.close()


def amount_convention(values):
    found = set()
    for raw in values:
        v = raw.strip()
        if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+,\d{1,2}", v, re.ASCII):
            found.add("point")
        if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+\.\d{1,2}", v, re.ASCII):
            found.add("comma")
    return "mixed" if len(found) > 1 else next(iter(found), "")


def read_amount(raw, convention):
    v = raw.strip()
    if convention == "mixed":
        raise ValueError("mixed")
    group = {"point": ".", "comma": ","}.get(convention)
    if "." in v and "," in v:
        inferred = amount_convention([v])
        if not inferred or (convention and convention != inferred):
            raise ValueError("shape")
        group = "." if inferred == "point" else ","
    elif re.fullmatch(r"-?\d{1,3}(?:[.,]\d{3})+", v, re.ASCII) and not group:
        raise ValueError("ambiguous")
    if group:
        decimal = "," if group == "." else "."
        pattern = rf"-?(?:\d+|\d{{1,3}}(?:{re.escape(group)}\d{{3}})+)(?:{re.escape(decimal)}\d{{1,2}})?"
    else:
        pattern = r"-?\d+(?:[.,]\d{1,2})?"
    if not re.fullmatch(pattern, v, re.ASCII):
        raise ValueError("shape")
    if group:
        v = v.replace(group, "")
    try:
        value = Decimal(v.replace(",", "."))
    except InvalidOperation as e:
        raise ValueError("shape") from e
    if not value.is_finite() or abs(value) > Decimal("999999999999.99"):
        raise ValueError("shape")
    return float(value)


def date_order(values):
    day = month = False
    for v in values:
        g = SLASHED_DATE.fullmatch(v.strip())
        if g:
            a, b, y = map(int, g.groups())
            try:
                datetime.date(y, b if a > 12 else a, a if a > 12 else b)
            except ValueError:
                continue
            day = day or a > 12
            month = month or b > 12
    return "mixed" if day and month else "dmy" if day else "mdy" if month else ""


def read_date(raw, order):
    v = raw.strip()
    if ISO_DATE.fullmatch(v):
        return datetime.date.fromisoformat(v).isoformat()
    g = SLASHED_DATE.fullmatch(v)
    if not g:
        raise ValueError("shape")
    if order not in {"dmy", "mdy"}:
        raise ValueError("mixed" if order == "mixed" else "ambiguous")
    a, b, y = map(int, g.groups())
    return datetime.date(y, b if order == "dmy" else a, a if order == "dmy" else b).isoformat()


def load_rows(st, rows, config, max_records=500, amount_format="auto", date_format="auto", status_map=None):
    pattern = config.records.get(config.actions["update_status"].record, "")
    convention = amount_convention([r.get("amount", "") for r in rows]) if amount_format == "auto" else amount_format
    order = date_order([r.get(f, "") for r in rows for f in ("issued", "due")]) if date_format == "auto" else date_format
    status_map = status_map or {}
    rejected, ids, skipped = [], [], []
    existing_ids = set(st.records.ids())
    for n, row in enumerate(rows, 2):
        rid, saw = row.get("id", "").strip(), ""
        try:
            if not rid or re.fullmatch(pattern, rid, re.ASCII) is None:
                raise ValueError(f"an id here must match the shape this adapter declares, {pattern}")
            if rid in existing_ids:
                skipped.append(rid)
                continue
            if len(existing_ids) >= max_records:
                raise ValueError("workspace record limit reached")
            customer = row.get("customer", "").strip()
            if not customer or len(customer) > 240:
                raise ValueError("customer must contain between 1 and 240 characters")
            currency = row.get("currency", "EUR").strip().upper() or "EUR"
            if not re.fullmatch(r"[A-Z]{3}", currency):
                raise ValueError("currency must be a three-letter code")
            status = row.get("status", "open").strip() or "open"
            status = status_map.get(status, status)
            if status not in STATUSES:
                raise ValueError("unknown status; declare a status mapping to open, paid, reminded, promised, disputed or cancelled")
            dates = {}
            for field in ("issued", "due"):
                saw = row.get(field, "").strip()
                try:
                    dates[field] = read_date(saw, order) if saw else ""
                except ValueError as e:
                    if str(e) == "mixed":
                        raise ValueError("this file's dates contradict each other; no row with these dates was loaded")
                    if str(e) == "ambiguous":
                        raise ValueError(f"{field} could be read day/month or month/day; declare the date format or use YYYY-MM-DD")
                    raise ValueError(f"{field} is not a date that exists: write it YYYY-MM-DD")
            saw = row.get("amount", "")
            try:
                amount = read_amount(saw, convention)
            except ValueError as e:
                if str(e) == "ambiguous":
                    raise ValueError("amount is ambiguous; a row written 1.234,56 or 1,234.56 would settle it for the whole file, or declare the format")
                if str(e) == "mixed":
                    raise ValueError("this file's amount formats contradict each other; declare one format")
                raise ValueError("amount is not a finite monetary value with valid grouping and at most two decimal places")
            contact = row.get("contact", "").strip()
            if contact and (len(contact) > 254 or not re.fullmatch(r"[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+", contact)):
                raise ValueError("contact must be one email address")
            # Import is a host operation. The complete accepted row is atomic, including its notes.
            with st.records.transaction():
                st.records.load_seed_row(dict(id=rid, customer=customer, amount=amount, currency=currency,
                                              status=status, contact=contact, **dates))
                if row.get("note"):
                    st.records.add_note_raw(rid, dates["issued"] or "uploaded", "uploaded", row["note"])
                if row.get("email_body"):
                    st.records.plant_email(rid, "uploaded", row.get("email_subject", ""), row["email_body"])
            ids.append(rid)
            existing_ids.add(rid)
        except ValueError as e:
            rejected.append(RejectedRow(row=n, id=rid, saw=saw, why=str(e)))
    tail = " — the row was not loaded, the rest were" if ids else " — the row was not loaded, and no row in this file was"
    for row in rejected:
        row.why += tail
    reading = []
    if order in {"dmy", "mdy"}:
        reading.append("dates read as " + ("day/month/year" if order == "dmy" else "month/day/year"))
    if convention in {"point", "comma"}:
        reading.append("amounts read with " + ("a point grouping and a comma deciding" if convention == "point" else "a comma grouping and a point deciding"))
    return UploadOut(loaded=len(ids), ids=ids, skipped=skipped, rejected=rejected,
                     read_as="; ".join(reading))
