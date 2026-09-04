"""The deployed face: a stranger's session, their own records, and the same boundary (PLAN.md §4.3).

Nothing here is a new way to write to records. A route can propose (through the graph), decide
(through `PolicyService.decide`, with a principal minted only by `api/auth.py`) and execute
(through `PolicyService.execute` with `agent/executor.py`) — the same three verbs the CLI has and
the tests have. There is no route that touches `records._apply_*`, and the grep test says so.

One session is one namespace: its own records database, its own policy store with its own audit
chain and its own fuse, its own checkpointer. Sessions cannot see each other, and a fuse tripped in
one binds only that one.
"""
from __future__ import annotations

import csv
import datetime
import io
import os
import re
import time
from pathlib import Path
from typing import Any

from fastapi import Body, FastAPI, Header, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse

from agent import build_graph
from agent.checkpoints import make_checkpointer
from agent.executor import make_executor
from agent.llm import GroqLLM, StubLLM
from api.auth import AGENT_PRINCIPAL, Sessions
from api.limits import Limits
from api.schemas import (ActionOut, AdapterOut, AssistOut, AuditOut, DecideIn, DecideOut, EmailOut,
                         FuseOut, NoteOut, ProposalOut, RecordDetailOut, RecordOut, RejectedRow,
                         SessionOut, UploadOut)
from policy import APPROVED, EXECUTED, HELD, PolicyConfig, PolicyService, Store
from records import Records

ROOT = Path(__file__).resolve().parents[1]
ADAPTER = os.environ.get("ATEZAIN_ADAPTER", "invoices-es")
STATE = Path(os.environ.get("ATEZAIN_STATE_DIR", ROOT / "var"))
# Render names the database URL DATABASE_URL; either works, and neither is ever written down here
DSN = os.environ.get("ATEZAIN_DSN") or os.environ.get("DATABASE_URL", "")
MODEL = os.environ.get("ATEZAIN_MODEL", "stub")
MODEL_ID = os.environ.get("ATEZAIN_MODEL_ID", "openai/gpt-oss-120b")
OWNER_TOKEN = os.environ.get("ATEZAIN_OWNER_TOKEN", "")

MAX_UPLOAD_BYTES = 1024 * 1024
MAX_ROWS = 500
COLUMNS = ("id", "customer", "amount", "currency", "issued", "due", "status")
EXTRAS = ("contact", "note", "email_subject", "email_body")

app = FastAPI(title="atezain", description=__doc__.split("\n\n")[0])
STATE.mkdir(parents=True, exist_ok=True)
sessions = Sessions(STATE / "sessions.db", dsn=DSN or None)   # the token table outlives a spin-down too
limits = Limits(state=STATE / "limits.json")
config = PolicyConfig.load(ROOT / "adapters" / ADAPTER / "permissions.toml")
_live: dict[str, "SessionState"] = {}
_health: dict[str, Any] = {"at": 0.0, "body": None}


def schema_of(sid: str) -> str:
    """One session, one Postgres schema. The id is hex from `secrets.token_hex`, but this does not
    take that on trust: anything that is not a lowercase word character is dropped before the name
    reaches a `CREATE SCHEMA` that cannot be parameterised."""
    return "s_" + "".join(ch for ch in sid.lower() if ch.isalnum() or ch == "_")[:48]


class SessionState:
    """One visitor's namespace: records, policy store, graph.

    With a DSN, all three are in Postgres and the records and the audit chain share ONE schema per
    session — which is what makes the deployed thing real: a Render free instance has no persistent
    disk and spins down when idle, so a visitor who comes back to their own link finds their
    invoices, their queue and their chain still there. Without a DSN it is SQLite files under
    `ATEZAIN_STATE_DIR`, which is the local shape and forgets nothing only because nothing spins
    down."""

    def __init__(self, sid: str):
        self.sid = sid
        if DSN:
            from policy.store_pg import PgStore
            from records.store_pg import PgRecords
            schema = schema_of(sid)
            self.records = PgRecords(DSN, schema=schema)
            self.policy = PolicyService(config, PgStore(DSN, schema=schema),
                                        record_reader=self.records.invoice)
            self.checkpointer = make_checkpointer(DSN, None)
        else:
            d = STATE / sid
            d.mkdir(parents=True, exist_ok=True)
            self.records = Records(str(d / "records.db"))
            self.policy = PolicyService(config, Store(str(d / "policy.db")),
                                        record_reader=self.records.invoice)
            self.checkpointer = make_checkpointer(None, d / "checkpoints.db")

    def graph(self, llm):
        return build_graph(self.records, self.policy, llm, AGENT_PRINCIPAL, ADAPTER, self.checkpointer)


def state_of(sid: str) -> SessionState:
    if sid not in _live:
        _live[sid] = SessionState(sid)
    return _live[sid]


def who(sid: str, authorization: str | None):
    """The session's human, or 401. The ONE mint lives in api/auth.py."""
    token = authorization.split(" ", 1)[1].strip() if authorization and " " in authorization else authorization
    p = sessions.principal(sid, token)
    if p is None:
        raise HTTPException(status_code=401, detail="no such session, or the token does not open it")
    return p


def model_for(byok: str | None):
    """The visitor's own key if they brought one, else the server's, else the stub."""
    if byok:
        return GroqLLM(model=MODEL_ID, api_key=byok), f"{MODEL_ID} (your key)", True
    if MODEL == "groq":
        return GroqLLM(model=MODEL_ID), MODEL_ID, False
    return StubLLM(), "stub", False


def out(p) -> ProposalOut:
    return ProposalOut(id=p.id, action=p.action, record_id=p.record_id, params=p.params,
                       status=p.status, reason=p.reason, decided_by=p.decided_by, note=p.note)


def proposals_of(st: SessionState) -> list:
    seen, rows = set(), []
    for r in st.policy.store.audit_rows():
        pid = r["proposal_id"]
        if r["kind"] != "PROPOSAL" or pid is None or pid in seen:
            continue
        seen.add(pid)
        p = st.policy.store.get_proposal(pid)
        if p is not None:
            rows.append(p)
    return rows


def answered(snap) -> bool:
    """A checkpoint is not an answer. LangGraph writes one for the thread as soon as the run
    starts, so a run that died inside `think` — the model refusing the call — leaves a checkpoint
    holding the input and nothing else, and `created_at is not None` then reads as *the model
    already answered, ask nobody again* for the life of the session. Client simulation 3
    (STATUS.md S3-5) lost seven of a visitor's records to one minute of a free tier's rate limit
    that way: 200, `cached`, empty summary, empty draft, no proposals, and no route that could ever
    ask again. `raw_proposals` is what `think` returns and nothing else writes; its presence is the
    model having spoken, and an answer with no proposals in it is still an answer."""
    return snap.created_at is not None and "raw_proposals" in snap.values


def model_refused(e: Exception, byok: bool) -> HTTPException:
    """The model saying no is not this service failing, and a stranger has to be able to tell the
    two apart — with their own key the difference is whether THEY have something to fix. Simulation
    3 (S3-4): a rate-limited call and a mistyped `X-Groq-Key` were both a bare 500."""
    code = getattr(e, "code", None)
    whose = "the key you sent" if byok else "the server's key"
    return HTTPException(status_code=502, detail=(
        f"the model did not answer ({code if code is not None else type(e).__name__}) using {whose}. "
        f"Nothing was written, and this record can be asked again."))


# ── sessions ─────────────────────────────────────────────────────────────────────────────────
@app.post("/sessions", response_model=SessionOut)
def create_session() -> SessionOut:
    sid, token = sessions.create()
    state_of(sid)
    return SessionOut(session=sid, token=token, adapter=ADAPTER,
                      note="the token is shown once; the server keeps only its sha256. "
                           "Send it as `Authorization: Bearer <token>`.")


@app.post("/sessions/{sid}/upload", response_model=UploadOut)
async def upload(sid: str, file: UploadFile, authorization: str | None = Header(default=None)) -> UploadOut:
    who(sid, authorization)
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"more than {MAX_UPLOAD_BYTES} bytes")
    rows = parse_rows(raw, file.filename or "")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=413, detail=f"more than {MAX_ROWS} rows")
    return load_rows(state_of(sid), rows)


def parse_rows(raw: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:                                    # pragma: no cover - openpyxl is installed
            raise HTTPException(status_code=415, detail="this deployment cannot read .xlsx; send CSV")
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(c).strip().lower() if c is not None else "" for c in next(it, [])]
        return [{k: ("" if v is None else str(v)) for k, v in zip(header, r)} for r in it]
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=415, detail="the file is not UTF-8 text or .xlsx")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "") for k, v in row.items()} for row in reader]


# ── how a file writes its numbers and dates, PROVED by the file and never guessed ────────────
# Client simulation 3 (STATUS.md S3-2, S3-3) refused a Spanish accounting export whole: `1.234,56`
# was "not a number" and `31/07/2026` was refused as a date. The open question was whether to READ
# them, and the reason it was open is that guessing is not allowed here — `1.234` is one thousand
# two hundred and thirty-four in one country and one point two three four in another, and
# `03/04/2026` is April in Bilbao and March in Boston. A wrong guess writes a number nobody typed
# into a record the assistant then reasons about.
#
# The way out is that a file usually settles its own convention, and where it does, nothing is
# being guessed: `1.234,56` carries both marks and the rightmost is the decimal one, and a `31` in
# the first position can only be a day. So the rule is: use a convention only where some row of
# THIS file proves it, apply it to the whole file, tell the visitor which convention was read and
# what proved it — and where the file proves nothing, refuse the row exactly as before and say
# that is why. Ruled 2026-09-04 (the judgment-dense model, on the owner's *do what you think is
# best*); the reasoning and the routes not taken are under STATUS.md Open questions.

ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
SLASHED_DATE = re.compile(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})")


def amount_convention(values) -> str:
    """`point` if this file writes 1.234,56 — the point groups and the comma decides — `comma` if
    it writes 1,234.56, and `""` if no row settles it. The proof is a value carrying BOTH marks:
    the rightmost of the two is the decimal mark, in every convention there is."""
    for v in values:
        i, j = v.rfind("."), v.rfind(",")
        if i >= 0 and j >= 0:
            return "point" if j > i else "comma"
    return ""


def read_amount(raw: str, convention: str) -> float:
    """The number, or ValueError('ambiguous') when this file has not settled which mark is which.
    Never a guess: `1.234` is refused unless some other row of the same file says what a point is."""
    v = raw.strip().replace(" ", "").replace("\u00a0", "")
    i, j = v.rfind("."), v.rfind(",")
    if i >= 0 and j >= 0:                                   # both marks: the rightmost decides
        dec, grp = (",", ".") if j > i else (".", ",")
        return float(v.replace(grp, "").replace(dec, "."))
    mark = "." if i >= 0 else ("," if j >= 0 else "")
    if not mark:
        return float(v)
    groups = {"point": ".", "comma": ","}.get(convention, "")
    if v.count(mark) > 1 or len(v.rsplit(mark, 1)[1]) == 3:  # 1.234.567, or 1.234 — group or decimal?
        if groups == mark:
            return float(v.replace(mark, ""))
        if groups:
            return float(v.replace(mark, "."))
        raise ValueError("ambiguous")
    return float(v.replace(mark, "."))                       # one mark, one or two figures after it


def date_order(values) -> str:
    """`dmy`, `mdy`, `mixed` when the file contradicts itself, or `""` when nothing settles it.
    The proof is a component over twelve: 31/07 can only be a day, 07/31 can only be a month."""
    day = month = False
    for v in values:
        g = SLASHED_DATE.fullmatch(v.strip())
        if g:
            day = day or int(g.group(1)) > 12
            month = month or int(g.group(2)) > 12
    if day and month:
        return "mixed"
    return "dmy" if day else ("mdy" if month else "")


def read_date(raw: str, order: str) -> str:
    """The date as YYYY-MM-DD, which is what the page compares and sorts. ValueError otherwise:
    `shape` for something that is not a date at all, `ambiguous` when the file has not settled the
    order, and whatever `datetime.date` says about 31/02."""
    v = raw.strip()
    if ISO_DATE.fullmatch(v):
        datetime.date.fromisoformat(v)
        return v
    g = SLASHED_DATE.fullmatch(v)
    if not g:
        raise ValueError("shape")
    if order not in ("dmy", "mdy"):
        raise ValueError("ambiguous" if order != "mixed" else "mixed")
    a, b, y = int(g.group(1)), int(g.group(2)), int(g.group(3))
    d, m = (a, b) if order == "dmy" else (b, a)
    return datetime.date(y, m, d).isoformat()


WHY_DATE_SHAPE = "{f} is not a date: write it YYYY-MM-DD, or the d/m/yyyy your own system exports"
WHY_DATE_AMBIGUOUS = ("{f} could be read day/month or month/day and nothing in this file settles which — one "
                      "date in it with a day past the twelfth would, and so would writing them YYYY-MM-DD")
WHY_DATE_MIXED = ("this file's dates contradict each other — some can only be day/month and others can only be "
                  "month/day — so no reading of the column is a safe one")
WHY_DATE_IMPOSSIBLE = "{f} is not a date that exists"
WHY_AMOUNT_AMBIGUOUS = ("amount could be a thousands separator or a decimal mark and nothing in this file settles "
                        "which — one row written 1.234,56 or 1,234.56 would settle it for the whole file")
WHY_AMOUNT_SHAPE = "amount is not a number"


def load_rows(st: SessionState, rows: list[dict]) -> UploadOut:
    """Every row is checked against the adapter before it becomes a record: an id the policy could
    never act on, or a date or an amount this file has not settled, is refused at the door with the
    reason instead of silently becoming a record whose every proposal is denied later — or worse,
    one the page then miscounts. Client simulation 3 (STATUS.md S3-2, S3-3) is why: `due` used to be
    stored as whatever string arrived, and the page answers *which of these do I chase* by comparing
    it to `YYYY-MM-DD`, so a `dd/mm/aaaa` export was sorted by the day of the month and its past-due
    count was wrong in both directions without a word."""
    pattern = config.records.get(config.actions["update_status"].record, "")
    convention = amount_convention([str(r.get("amount", "")) for r in rows])
    order = date_order([str(r.get(f, "")) for r in rows for f in ("issued", "due")])
    rejected, loaded, ids = [], 0, []
    for n, row in enumerate(rows, start=2):                     # row 1 is the header
        rid = (row.get("id") or "").strip()
        if not rid:
            rejected.append(RejectedRow(row=n, id="", why="no id"))
            continue
        if re.fullmatch(pattern, rid, re.ASCII) is None:
            rejected.append(RejectedRow(row=n, id=rid, why=f"an id here must match the shape this adapter declares, {pattern}"))
            continue
        dates, bad = {}, None
        for f in ("issued", "due"):
            v = (row.get(f) or "").strip()
            if not v:
                dates[f] = ""
                continue
            try:
                dates[f] = read_date(v, order)
            except ValueError as e:
                why = {"shape": WHY_DATE_SHAPE, "ambiguous": WHY_DATE_AMBIGUOUS,
                       "mixed": WHY_DATE_MIXED}.get(str(e), WHY_DATE_IMPOSSIBLE)
                bad = RejectedRow(row=n, id=rid, saw=v, why=why.format(f=f))
                break
        if bad is not None:
            rejected.append(bad)
            continue
        existing = st.records.invoice(rid)
        if existing is None:
            raw = str(row.get("amount", "0")).strip()
            try:
                amount = read_amount(raw, convention)
            except ValueError as e:
                rejected.append(RejectedRow(row=n, id=rid, saw=raw,
                                            why=WHY_AMOUNT_AMBIGUOUS if str(e) == "ambiguous" else WHY_AMOUNT_SHAPE))
                continue
            st.records.load_seed_row({
                "id": rid, "customer": (row.get("customer") or "").strip(), "amount": amount,
                "currency": (row.get("currency") or "EUR").strip(),
                "issued": dates["issued"], "due": dates["due"],
                "status": (row.get("status") or "open").strip(),
                "contact": (row.get("contact") or "").strip()})
            loaded += 1
            ids.append(rid)
        if row.get("note"):
            st.records.add_note_raw(rid, dates.get("issued") or "uploaded", "uploaded", row["note"])
        if row.get("email_body"):
            st.records.plant_email(rid, "uploaded", row.get("email_subject", ""), row["email_body"])
    # What happened to the OTHER rows is not something a per-row reason can know, and saying it
    # anyway made a file where every single row was refused report, five hundred times over, that
    # the rest had loaded (client simulation 3, STATUS.md S3-1).
    tail = (" — the row was not loaded, the rest were" if loaded
            else " — the row was not loaded, and no row in this file was")
    rejected = [RejectedRow(row=r.row, id=r.id, saw=r.saw, why=r.why + tail) for r in rejected]
    # What was read, and what proved it. A layer that reinterprets somebody's money owes them the
    # sentence saying how (client simulation 3, STATUS.md S3-2).
    read = []
    if order in ("dmy", "mdy"):
        read.append(f"dates read as {'day/month/year' if order == 'dmy' else 'month/day/year'}, "
                    f"which this file's own dates settle")
    if convention:
        read.append(f"amounts read with {'a point grouping and a comma deciding' if convention == 'point' else 'a comma grouping and a point deciding'}, "
                    f"which a row of this file carrying both marks settles")
    return UploadOut(loaded=loaded, rejected=rejected, ids=ids, read_as="; ".join(read))


@app.get("/adapter", response_model=AdapterOut)
def adapter() -> AdapterOut:
    """The permission table, rendered from the loaded config — not from a copy of it in the page.
    No session and no token: it is the same for everyone, and someone deciding whether to trust the
    layer should be able to read the rules before uploading anything to it."""
    return AdapterOut(
        adapter=config.adapter, version=config.version, fingerprint=config.fingerprint(),
        daily_writes=config.daily_writes, records=dict(config.records),
        actions=[ActionOut(action=name, record=s.record, writes=list(s.writes), approval=s.approval,
                           denied=s.deny, daily_max=s.daily_max,
                           values={k: list(v) for k, v in s.constraints.items()},
                           of_the_record=dict(s.record_constraints))
                 for name, s in sorted(config.actions.items())])


def record_out(inv: dict) -> RecordOut:
    return RecordOut(id=inv["id"], customer=inv["customer"], amount=inv["amount"], currency=inv["currency"],
                     issued=inv["issued"], due=inv["due"], status=inv["status"],
                     contact=inv.get("contact"), reminder_to=inv.get("reminder_to"),
                     reminder_channel=inv.get("reminder_channel"),
                     notes=len(inv["notes"]), emails=len(inv["emails"]),
                     assistant_notes=sum(1 for n in inv["notes"] if n["author"] == "assistant"))


@app.get("/sessions/{sid}/records", response_model=list[RecordOut])
def records(sid: str, authorization: str | None = Header(default=None)) -> list[RecordOut]:
    """The visitor's own records back, which until 2026-09-03 no route returned: an `add_note` is
    auto-approved by this adapter and executes without anyone deciding, so the one write that needs
    no human was the one write nobody could read (client simulation 1, STATUS.md). Reads only —
    there is no route here that writes."""
    who(sid, authorization)
    st = state_of(sid)
    return [record_out(inv) for inv in (st.records.invoice(i) for i in st.records.ids()) if inv]


@app.get("/sessions/{sid}/records/{invoice_id}", response_model=RecordDetailOut)
def record(sid: str, invoice_id: str, authorization: str | None = Header(default=None)) -> RecordDetailOut:
    """One record with its notes and emails. Every note carries WHO wrote it: `assistant` is this
    system's own voice, and a claim in one is the assistant's, not the customer's."""
    who(sid, authorization)
    inv = state_of(sid).records.invoice(invoice_id)
    if inv is None:
        raise HTTPException(status_code=404, detail=f"no record {invoice_id} in this session")
    return RecordDetailOut(**record_out(inv).model_dump(),
                           note_rows=[NoteOut(**n) for n in inv["notes"]],
                           email_rows=[EmailOut(**e) for e in inv["emails"]])


# ── assist ───────────────────────────────────────────────────────────────────────────────────
@app.post("/sessions/{sid}/assist/{invoice_id}", response_model=AssistOut)
def assist(sid: str, invoice_id: str, request: Request, authorization: str | None = Header(default=None),
           x_groq_key: str | None = Header(default=None)) -> AssistOut:
    who(sid, authorization)
    ok, why = limits.allow(request.client.host if request.client else "?")
    if not ok:
        # WHICH limit, and that their own key is not the answer to it. A visitor who brought one has
        # been told they are paying their own way, and is then stopped by a limit that is about this
        # instance and not about the model — `rate_limited_minute` alone said neither (client
        # simulation 3, STATUS.md S3-7)
        raise HTTPException(status_code=429, detail=(
            f"{why}: one address gets {limits.per_minute} assists a minute and {limits.per_day} a day "
            f"on this instance. That limit is this instance's, not the model's — your own "
            f"`X-Groq-Key` is neither counted against the server's model budget nor able to lift it."))
    st = state_of(sid)
    if st.records.invoice(invoice_id) is None:
        raise HTTPException(status_code=404, detail=f"no record {invoice_id} in this session")

    llm, model_name, byok = model_for(x_groq_key)
    conf = {"configurable": {"thread_id": f"{sid}:{invoice_id}"}}
    graph = st.graph(llm)
    snap = graph.get_state(conf)
    if answered(snap):
        # the model already answered for this record: give back what it said, ask nobody again
        s = snap.values
        return AssistOut(invoice_id=invoice_id, summary=s.get("summary", ""),
                         recommendation=s.get("recommendation", ""), draft=s.get("draft", ""),
                         proposals=[out(p) for p in proposals_of(st) if p.record_id == invoice_id],
                         cached=True, model=model_name)

    allowed, why = limits.allow_model_call(byok)
    if not allowed:
        raise HTTPException(status_code=503, detail=why)
    try:
        s = graph.invoke({"invoice_id": invoice_id, "task": "draft"}, config=conf)
    except HTTPException:
        raise
    except Exception as e:                                      # noqa: BLE001 — see model_refused
        raise model_refused(e, byok) from e
    return AssistOut(invoice_id=invoice_id, summary=s.get("summary", ""),
                     recommendation=s.get("recommendation", ""), draft=s.get("draft", ""),
                     proposals=[out(p) for p in proposals_of(st) if p.record_id == invoice_id],
                     cached=False, model=model_name)


# ── the queue, the decision, the chain ───────────────────────────────────────────────────────
@app.get("/sessions/{sid}/proposals", response_model=list[ProposalOut])
def queue(sid: str, authorization: str | None = Header(default=None)) -> list[ProposalOut]:
    who(sid, authorization)
    return [out(p) for p in proposals_of(state_of(sid))]


@app.post("/sessions/{sid}/proposals/{pid}/decide", response_model=DecideOut)
def decide(sid: str, pid: str, body: DecideIn = Body(...), authorization: str | None = Header(default=None)) -> DecideOut:
    human = who(sid, authorization)
    st = state_of(sid)
    if st.policy.store.get_proposal(pid) is None:
        raise HTTPException(status_code=404, detail="no such proposal in this session")
    p = st.policy.decide(pid, body.approve, human, note=body.note)
    if p.status != APPROVED:
        return DecideOut(proposal=out(p), executed=False)
    q = st.policy.execute(p.id, make_executor(st.records), human)
    effect = st.policy.store.get_proposal(q.id)
    return DecideOut(proposal=out(effect or q), executed=q.status == EXECUTED,
                     applied=q.params if q.status == EXECUTED else None)


@app.get("/sessions/{sid}/audit", response_model=AuditOut)
def audit(sid: str, authorization: str | None = Header(default=None)) -> AuditOut:
    who(sid, authorization)
    st = state_of(sid)
    seq, h = st.policy.store.audit_head()
    return AuditOut(rows=st.policy.store.audit_rows(), head_seq=seq, head_hash=h,
                    verifies=st.policy.store.audit_verify(),
                    anomalies=[list(a) for a in st.policy.store.audit_anomalies()],
                    fuse=st.policy.store.fuse_get())


@app.post("/sessions/{sid}/fuse/clear", response_model=FuseOut)
def clear_fuse(sid: str, authorization: str | None = Header(default=None)) -> FuseOut:
    human = who(sid, authorization)
    st = state_of(sid)
    cleared = st.policy.fuse.clear(human)
    return FuseOut(cleared=cleared, fuse=st.policy.store.fuse_get(),
                   why="" if cleared else "a tripped fuse clears no earlier than the day after it tripped")


@app.post("/server/fuse/clear")
def clear_server_fuse(x_owner_token: str | None = Header(default=None)) -> JSONResponse:
    """The SERVER's model budget, not a session's fuse. The owner's token, or nothing."""
    ok = limits.clear_fuse(x_owner_token or "", OWNER_TOKEN or None)
    return JSONResponse({"cleared": ok, "model_calls": limits.model_calls, "fuse": limits.fuse},
                        status_code=200 if ok else 403)


# ── the page, and the health check ───────────────────────────────────────────────────────────
@app.get("/")
def root() -> RedirectResponse:
    """A link shared without the path landed on a bare 404 (client simulation 2). What a root page
    should SAY is the owner's; that it should not be an error is not a question."""
    return RedirectResponse(url="/demo", status_code=307)


@app.get("/demo", response_class=HTMLResponse)
def demo() -> str:
    return (Path(__file__).resolve().parent / "demo.html").read_text(encoding="utf-8")


@app.get("/healthz")
def healthz() -> dict:
    now = time.time()
    if _health["body"] is not None and now - _health["at"] < 60:
        return _health["body"]
    try:
        sessions.exists("healthz")
        store_ok = True
    except Exception:                                          # noqa: BLE001
        store_ok = False
    body = {"store": store_ok, "backing": "postgres" if DSN else "sqlite",
            "model": MODEL if MODEL != "groq" else MODEL_ID,
            "model_calls_today": limits.model_calls, "server_fuse": limits.fuse,
            "adapter": ADAPTER, "sessions": len(sessions.ids())}
    _health.update({"at": now, "body": body})
    return body
